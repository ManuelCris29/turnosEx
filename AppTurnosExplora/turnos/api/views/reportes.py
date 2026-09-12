import logging
import math

from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import JsonResponse
from django.views.generic import View

from core.mixins import SupervisorApiRequiredMixin
from core.services.cache_service import CACHE_TTL_SHORT, CacheService
from core.utils.date_utils import DateUtils


# ─────────────────────────────────────────────────────────────────────────────
# Helpers de texto "a prueba de bobos": convierten los datos crudos del reporte
# en frases explicativas claras. Reutilizados por las 4 hojas del Excel.
# ─────────────────────────────────────────────────────────────────────────────
def _fmt_fecha(iso):
    """'2026-07-15' -> '15/07/2026'. None/'' -> ''."""
    if not iso:
        return ''
    try:
        y, m, d = str(iso).split('-')
        return f'{d}/{m}/{y}'
    except Exception:
        return str(iso)


def _frase_trabaja(emp, dia_info):
    """Explica POR QUÉ esta persona trabaja hoy (y a quién cubre)."""
    tipo = emp.get('tipo', 'oficial')
    jd = emp.get('jornada_dia', '') or ''
    cubre = (emp.get('cubre_a') or {}).get('nombre', '')
    deuda = emp.get('deuda_reprogramacion') or {}

    if deuda.get('paga_hoy'):
        return (f'Dobla hoy (AM y PM) para pagar la doblada que debía '
                f'del {_fmt_fecha(deuda.get("fecha_original"))}.')
    if tipo == 'doblada':
        if jd == 'DOBLADA':
            return (f'Dobla hoy: trabaja AM y PM para cubrir a {cubre}, que descansa.'
                    if cubre else 'Dobla hoy: trabaja AM y PM cubriendo a un compañero que descansa.')
        return (f'Trabaja también la jornada {jd} cubriendo a {cubre}, además de la suya.'
                if cubre else f'Trabaja también la jornada {jd} cubriendo a un compañero.')
    if tipo == 'cambio':
        if jd == 'DOBLADA':
            return 'Dobla hoy (AM y PM) por un cambio de turno aprobado.'
        return f'Trabaja hoy en {jd} por un cambio de turno aprobado.'
    # oficial
    if jd == 'DOBLADA':
        rot = ('el fin de semana' if dia_info.get('es_finde')
               else 'un festivo' if dia_info.get('es_festivo') else 'rotación')
        return f'Le toca doblar por rotación ({rot}): trabaja AM y PM.'
    return f'Turno normal: trabaja su jornada {jd} de siempre.'


def _frase_descansa(emp):
    """Explica POR QUÉ esta persona descansa hoy (y con quién se relaciona)."""
    motivo = (emp.get('motivo') or '').lower()
    comp = (emp.get('companero') or {}).get('nombre', '')
    jb = emp.get('jornada_base') or ''
    frases = {
        'cedió su jornada': (f'Descansa porque cedió su jornada de hoy a {comp}, que la '
                             f'trabaja por él/ella (se la devolverá otro día).'),
        'paga doblada': (f'Descansa porque hoy le devuelven la doblada: {comp} trabaja '
                         f'doble y cubre su jornada.'),
        'cambio de día de descanso': f'Descansa hoy por un cambio de día de descanso acordado con {comp}.',
        'doblada permanente': (f'Descansa por acuerdo de doblada permanente: {comp} cubre su '
                               f'jornada este día de la semana.'),
        'festivo': 'Descansa porque hoy es festivo y no le toca doblar por rotación.',
        'descanso de fin de semana': f'Descansa porque este fin de semana no le toca a su grupo ({jb}).',
        'descanso de temporada': f'Descansa por descanso de temporada asignado a su grupo ({jb}).',
        'lunes de mantenimiento': 'Descansa porque hoy es lunes de mantenimiento (no hay operación).',
        'sin jornada asignada': 'No tiene jornada asignada actualmente (verificar con el supervisor).',
        'sin alternancia publicada': ('⚠ SIN PLANIFICAR: este día no tiene publicada la alternancia '
                                      'de findes/festivos, así que NO se sabe si trabaja. No es un '
                                      'descanso: hay que publicar la alternancia del año.'),
    }
    return frases.get(motivo, (emp.get('motivo') or 'Descansa').capitalize())


#: Cómo se lee el papel de la persona en el trato. El vocabulario interno
#: ('solicitante'/'receptor') no le dice nada a un supervisor mirando una planilla.
ROL_LBL = {'solicitante': 'Pidió el cambio', 'receptor': 'Recibió el cambio'}


def _celdas_acuerdo(emp):
    """Las 5 celdas del ACUERDO que puso a esta persona en la jornada de hoy.

    Mismo dict en las dos hojas (`ReporteDiaService.reporte` lo normaliza), así que hay un
    solo formateador: con quién, de qué tipo, qué papel juega, cuál es el OTRO día del trato
    y cuándo se aprobó. Todo '—' cuando el día no viene de ninguna solicitud (turno normal,
    rotación, festivo, temporada…).
    """
    a = emp.get('acuerdo') or {}
    if not a:
        return ['—'] * 5
    return [
        a.get('companero_nombre') or '—',
        a.get('tipo') or '—',
        ROL_LBL.get(a.get('rol')) or '—',
        a.get('fecha_relacionada') or '—',
        a.get('fecha_resolucion') or '—',
    ]


def _txt_permiso(emp):
    p = emp.get('permiso')
    if not p:
        return ''
    horas = p.get('horas', 0) or 0
    esp = p.get('especificacion', '')
    esp = f': {esp}' if esp else ''
    return f'🕐 Permiso {horas:g}h — {p.get("tipo","")}{esp} ({p.get("estado","")})'


def _txt_restriccion(emp):
    r = emp.get('restriccion')
    if not r:
        return ''
    hasta = _fmt_fecha(r.get('fecha_fin')) or 'indefinida'
    rec = f' — {r["recomendacion"]}' if r.get('recomendacion') else ''
    return f'⚠ Restricción: {r.get("tipo","")}{rec} (hasta {hasta})'


def _txt_sancion(emp):
    s = emp.get('sancion')
    if not s:
        return ''
    hasta = _fmt_fecha(s.get('fecha_fin')) or 'indefinida'
    return f'⛔ Sancionado/a hasta {hasta}: {s.get("motivo","")}'


def _txt_deuda(emp):
    d = emp.get('deuda_reprogramacion')
    if not d:
        return ''
    orig = _fmt_fecha(d.get('fecha_original'))
    if d.get('paga_hoy'):
        return f'💥 HOY paga su doblada reprogramada (debía la del {orig})'
    if d.get('fecha_reprogramada'):
        return f'🔁 Debe una doblada del {orig}; la pagará el {_fmt_fecha(d["fecha_reprogramada"])}'
    jd = d.get('jornada_debida') or ''
    jd = f' ({jd})' if jd else ''
    return f'🔁 Debe una doblada{jd} del {orig} — aún sin fecha de pago'


def _flags_cortos(emp):
    """Versión corta de los flags para la hoja Resumen."""
    out = []
    p = emp.get('permiso')
    if p:
        out.append(f'🕐 {(p.get("horas",0) or 0):g}h permiso')
    if emp.get('restriccion'):
        out.append('⚠ Restricción')
    if emp.get('sancion'):
        out.append('⛔ Sanción')
    d = emp.get('deuda_reprogramacion')
    if d:
        out.append('💥 Hoy paga doblada' if d.get('paga_hoy')
                   else f'🔁 Debe doblada del {_fmt_fecha(d.get("fecha_original"))}')
    return out


def _altura(*pairs):
    """Altura de fila según el texto más largo respecto al ancho de su columna."""
    lineas = 1
    for texto, ancho in pairs:
        if texto:
            lineas = max(lineas, math.ceil(len(str(texto)) / max(8, ancho - 2)))
    return min(140, max(16, lineas * 14 + 2))


MESES_ES = ['', 'enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio',
            'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']
DIAS_ES = ['lunes', 'martes', 'miércoles', 'jueves', 'viernes', 'sábado', 'domingo']


def _deuda_del_corte(fecha):
    """
    Deuda todavía sin pagar del día 1 del mes de `fecha` hasta `fecha`, y su agregado.

    Misma fuente de verdad que /empleados/sanciones/morosos/?corte=…: la regla vive en
    `DeudaCorporativaService.deuda_a_corte` y aquí solo se consulta. La comparten la vista
    JSON (que expone únicamente el agregado) y la de Excel (que lleva además el detalle),
    para que las dos no puedan contar cosas distintas del mismo día.

    Devuelve (filas, resumen) con resumen = {exploradores, minutos, horas, proyectada}.
    `proyectada` avisa de que la fecha es futura: entonces la cifra incluye días que
    todavía no han vencido, así que no es deuda exigible, es una previsión.

    Se cachean 5 minutos las `filas` (la consulta cara, determinística en función
    de `fecha`), igual razón que `ReporteDiaService.reporte`: hoy se calculan por
    separado desde `ReporteDiaView` y `ReporteDiaExcelView` para la MISMA fecha,
    y no hay un evento único de invalidación (depende de aprobaciones, sanciones
    y pagos de deuda). `proyectada` se recalcula SIEMPRE fuera de la caché: no
    depende de `fecha` sola sino de "hoy" (`timezone.localdate()`), que cambia
    aunque `fecha` no cambie — cachearla junto con `filas` serviría un rótulo
    PROYECTADA/PENDIENTE desactualizado durante todo el TTL.
    """
    from django.utils import timezone

    from solicitudes.services.deuda_corporativa_service import DeudaCorporativaService

    filas = CacheService.get_or_set(
        f"deuda_corte_filas_v1_{fecha}",
        lambda: DeudaCorporativaService.deuda_a_corte(fecha),
        ttl=CACHE_TTL_SHORT,
    )
    minutos = sum(f['minutos'] for f in filas)
    return filas, {
        'exploradores': len(filas),
        'minutos': minutos,
        # En horas, que es como se paga: el PDH se registra en horas.
        'horas': round(minutos / 60, 2),
        'proyectada': fecha > timezone.localdate(),
    }


def _fecha_de_request(request):
    """
    Lee y valida ?fecha=YYYY-MM-DD.

    Devuelve (fecha, None) si es válida, o (None, JsonResponse de error) si no.
    """
    fecha_str = request.GET.get('fecha')
    if not fecha_str:
        return None, JsonResponse({'error': 'Debe enviar fecha (YYYY-MM-DD)'}, status=400)
    try:
        return DateUtils.parse_date(fecha_str), None
    except (ValueError, TypeError):
        return None, JsonResponse({'error': 'Formato de fecha inválido'}, status=400)


def _error_500(nombre_vista, exc):
    """
    Registra el detalle en el log y devuelve un mensaje genérico.

    El texto de la excepción puede contener rutas, SQL o nombres de tabla: no se
    manda al navegador, que además lo pintaría dentro del HTML del reporte.
    """
    logging.getLogger(__name__).error('%s: %s', nombre_vista, exc, exc_info=True)
    return JsonResponse(
        {'error': 'No se pudo generar el reporte. Inténtalo de nuevo o avisa al administrador.'},
        status=500,
    )


class ReporteDiaView(LoginRequiredMixin, SupervisorApiRequiredMixin, View):
    """Reporte operacional del día para supervisores. Solo accesible con rol supervisor/staff."""

    def get(self, request):
        fecha, error = _fecha_de_request(request)
        if error:
            return error

        try:
            from turnos.services.reporte_dia_service import ReporteDiaService
            data = ReporteDiaService.reporte(fecha)
            # Solo el agregado: la pantalla es un panorama de la operación y el detalle de
            # quién debe ya tiene la suya (/empleados/sanciones/morosos/). El Excel sí lo lleva.
            data['deuda_corte'] = _deuda_del_corte(fecha)[1]
            return JsonResponse(data)
        except Exception as e:
            return _error_500('ReporteDiaView', e)


class ReporteMesDiasView(LoginRequiredMixin, SupervisorApiRequiredMixin, View):
    """
    Tipo de día (festivo / finde / mantenimiento / sin planificar) de todo un mes.

    Alimenta el pintado del calendario: una petición por mes en vez de una por celda.
    Es solo el calendario, sin datos de personas, así que es barato.
    """

    def get(self, request):
        try:
            anio = int(request.GET.get('anio', ''))
            mes = int(request.GET.get('mes', ''))
        except (TypeError, ValueError):
            return JsonResponse({'error': 'Debe enviar anio y mes numéricos'}, status=400)
        if not 1 <= mes <= 12:
            return JsonResponse({'error': 'El mes debe estar entre 1 y 12'}, status=400)

        try:
            from turnos.services.reporte_dia_service import ReporteDiaService
            return JsonResponse({'dias': ReporteDiaService.dias_del_mes(anio, mes)})
        except Exception as e:
            return _error_500('ReporteMesDiasView', e)


class ReporteDiaExcelView(LoginRequiredMixin, SupervisorApiRequiredMixin, View):
    """Exporta el reporte operacional del día como archivo Excel (.xlsx)."""

    def get(self, request):
        fecha, error = _fecha_de_request(request)
        if error:
            return error

        try:
            from turnos.services.reporte_dia_service import ReporteDiaService
            data = ReporteDiaService.reporte(fecha)
            # Se consulta aquí, no dentro del generador, para que el Excel siga siendo
            # solo formato: `_generar_excel` recibe datos ya resueltos.
            filas_deuda, resumen_deuda = _deuda_del_corte(fecha)
            return self._generar_excel(fecha, data, filas_deuda, resumen_deuda)
        except Exception as e:
            return _error_500('ReporteDiaExcelView', e)

    def _generar_excel(self, fecha, data, filas_deuda=None, resumen_deuda=None):
        import io

        from django.http import HttpResponse
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        fecha_legible = (f'{DIAS_ES[fecha.weekday()]} {fecha.day} de '
                         f'{MESES_ES[fecha.month]} de {fecha.year}')

        trabajando  = data.get('trabajando', [])
        descansando = data.get('descansando', [])
        dia_info    = data.get('dia_info', {})

        am   = [e for e in trabajando if e.get('jornada_dia') in ('AM', 'DOBLADA')]
        pm   = [e for e in trabajando if e.get('jornada_dia') in ('PM', 'DOBLADA')]
        desc = descansando

        # ── Estilos ──────────────────────────────────────────────────────────
        def fill(hex_color):
            return PatternFill('solid', fgColor=hex_color)

        AZUL_OSCURO  = '1E3A5F'
        AZUL_MEDIO   = '2563EB'
        AZUL_CLARO   = 'DBEAFE'
        NARANJA_OSC  = '92400E'
        NARANJA_MED  = 'D97706'
        NARANJA_CL   = 'FEF3C7'
        GRIS_OSC     = '334155'
        GRIS_MED     = '64748B'
        GRIS_CLARO   = 'F1F5F9'
        VERDE_CL     = 'DCFCE7'
        MORADO_CL    = 'EDE9FE'
        ROJO_OSC     = '7F1D1D'
        ROJO_MED     = 'DC2626'
        BLANCO       = 'FFFFFF'

        def fuente(bold=False, color='000000', size=10, italic=False):
            return Font(bold=bold, color=color, size=size, italic=italic)

        def borde_fino():
            s = Side(style='thin', color='CBD5E1')
            return Border(left=s, right=s, top=s, bottom=s)

        def centrado(wrap=False):
            return Alignment(horizontal='center', vertical='center', wrap_text=wrap)

        def izquierda(wrap=False):
            return Alignment(horizontal='left', vertical='center', wrap_text=wrap)

        wb = Workbook()
        wb.remove(wb.active)  # quitar hoja vacía por defecto

        # ════════════════════════════════════════════════════════════════════
        # HOJA 1 — RESUMEN
        # ════════════════════════════════════════════════════════════════════
        ws = wb.create_sheet('Resumen')
        ws.sheet_view.showGridLines = False
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 18
        ws.column_dimensions['C'].width = 34
        ws.column_dimensions['D'].width = 34

        # Título principal
        ws.merge_cells('A1:D1')
        c = ws['A1']
        c.value = '📋  REPORTE OPERACIONAL'
        c.font = fuente(bold=True, color=BLANCO, size=14)
        c.fill = fill(AZUL_OSCURO)
        c.alignment = centrado()
        ws.row_dimensions[1].height = 32

        # Fecha
        ws.merge_cells('A2:D2')
        c = ws['A2']
        c.value = fecha_legible.capitalize()
        c.font = fuente(bold=False, color=BLANCO, size=11, italic=True)
        c.fill = fill(AZUL_MEDIO)
        c.alignment = centrado()
        ws.row_dimensions[2].height = 22

        # Tipo de día
        tipo_dia = ('🎉 Festivo' if dia_info.get('es_festivo')
                    else '🏖️ Fin de semana' if dia_info.get('es_finde')
                    else '🔧 Mantenimiento' if dia_info.get('es_mantenimiento')
                    else '💼 Día laboral')
        if dia_info.get('sin_planificar'):
            tipo_dia += '  ·  ⚠ SIN ALTERNANCIA PUBLICADA: este día no está planificado'
        ws.merge_cells('A3:D3')
        c = ws['A3']
        c.value = tipo_dia
        c.font = fuente(bold=True, size=10, color=('991B1B' if dia_info.get('sin_planificar') else '000000'))
        c.fill = fill('FEE2E2' if dia_info.get('sin_planificar') else 'EFF6FF')
        c.alignment = centrado()
        ws.row_dimensions[3].height = 18

        ws.row_dimensions[4].height = 8  # separador

        # Contadores. La 1ª columna es el total de empleados activos (todos los que
        # aparecen en el reporte, trabajando o descansando); las otras 3 son el desglose.
        headers_res = ['Total empleados', 'Trabajan AM', 'Trabajan PM', 'Descansan']
        colores_res = [GRIS_CLARO, AZUL_CLARO, NARANJA_CL, GRIS_CLARO]
        fuentes_res = [GRIS_OSC, AZUL_MEDIO, NARANJA_OSC, GRIS_OSC]
        valores_res = [len(trabajando) + len(descansando), len(am), len(pm), len(desc)]

        for col, (h, bg, fg, v) in enumerate(
                zip(headers_res, colores_res, fuentes_res, valores_res), 1):
            # Etiqueta
            lbl = ws.cell(row=5, column=col, value=h)
            lbl.font = fuente(bold=True, color=fg, size=10)
            lbl.fill = fill(bg)
            lbl.alignment = centrado()
            lbl.border = borde_fino()
            ws.row_dimensions[5].height = 20
            # Número
            num = ws.cell(row=6, column=col, value=v)
            num.font = fuente(bold=True, color=fg, size=16)
            num.fill = fill(bg)
            num.alignment = centrado()
            num.border = borde_fino()
            ws.row_dimensions[6].height = 36

        # ── Deuda pendiente del mes hasta esta fecha ─────────────────────────
        # El detalle está en la hoja "Deuda pendiente"; aquí va el titular, para que no
        # haya que cambiar de hoja para enterarse de que hay gente debiendo.
        rd = resumen_deuda or {}
        ws.merge_cells('A7:D7')
        c = ws['A7']
        if not rd.get('exploradores'):
            c.value = f'✅ Nadie debe horas del 1 al {fecha.day} de {MESES_ES[fecha.month]}'
            c.font = fuente(bold=True, color='166534', size=10)
            c.fill = fill('DCFCE7')
        else:
            etiqueta = 'Deuda PROYECTADA' if rd.get('proyectada') else 'Deuda pendiente'
            c.value = (f'💸  {etiqueta} del 1 al {fecha.day} de {MESES_ES[fecha.month]}: '
                       f'{rd["exploradores"]} explorador(es) · {rd["horas"]} h '
                       f'({rd["minutos"]} min)  ·  detalle en la hoja "Deuda pendiente"')
            c.font = fuente(bold=True, color='991B1B', size=10)
            c.fill = fill('FEE2E2')
        c.alignment = centrado()
        c.border = borde_fino()
        ws.row_dimensions[7].height = 20

        # Mini-lista resumen: por persona, etiqueta corta + frase explicativa + flags
        TIPO_BG = {'oficial': 'DBEAFE', 'cambio': 'FEF9C3', 'doblada': 'EDE9FE'}
        TIPO_FG = {'oficial': '1D4ED8', 'cambio': '92400E', 'doblada': '5B21B6'}
        TIPO_LBL = {'oficial': 'Turno normal', 'cambio': 'Cambio', 'doblada': 'Dobla'}

        fila = 8
        for seccion, empleados_sec, bg_h, fg_h, trabaja in [
            ('👔  TRABAJAN AM', am,   AZUL_CLARO,  AZUL_MEDIO,  True),
            ('🌙  TRABAJAN PM', pm,   NARANJA_CL,  NARANJA_MED, True),
            ('🛌  DESCANSAN',   desc, GRIS_CLARO,  GRIS_MED,    False),
        ]:
            ws.merge_cells(f'A{fila}:D{fila}')
            c = ws.cell(row=fila, column=1, value=seccion)
            c.font = fuente(bold=True, color=fg_h, size=10)
            c.fill = fill(bg_h)
            c.alignment = izquierda()
            c.border = borde_fino()
            ws.row_dimensions[fila].height = 18
            fila += 1
            for emp in empleados_sec:
                nombre = f'{emp["nombre"]} {emp["apellido"]}'
                deuda = emp.get('deuda_reprogramacion') or {}
                if trabaja:
                    tipo = emp.get('tipo', 'oficial')
                    etiqueta = 'HOY PAGA DOBLADA' if deuda.get('paga_hoy') else TIPO_LBL.get(tipo, 'Trabaja')
                    et_bg = 'DCFCE7' if deuda.get('paga_hoy') else TIPO_BG.get(tipo, 'FFFFFF')
                    et_fg = '166534' if deuda.get('paga_hoy') else TIPO_FG.get(tipo, '000000')
                    frase = _frase_trabaja(emp, dia_info)
                else:
                    etiqueta = 'Descansa'
                    et_bg, et_fg = 'F1F5F9', GRIS_OSC
                    frase = _frase_descansa(emp)
                detalle = ' · '.join([frase] + _flags_cortos(emp))

                # Col A: nombre (fill de alerta si sanción / restricción)
                c = ws.cell(row=fila, column=1, value=nombre)
                nombre_bg = ('FEE2E2' if emp.get('sancion')
                             else 'FEF3C7' if emp.get('restriccion') else 'FFFFFF')
                c.font = fuente(bold=True, size=10)
                c.fill = fill(nombre_bg)
                c.alignment = izquierda()
                c.border = borde_fino()
                # Col B: etiqueta corta de situación
                cb = ws.cell(row=fila, column=2, value=etiqueta)
                cb.font = fuente(bold=True, color=et_fg, size=9)
                cb.fill = fill(et_bg)
                cb.alignment = centrado(True)
                cb.border = borde_fino()
                # Col C:D: frase + flags
                ws.merge_cells(f'C{fila}:D{fila}')
                c2 = ws.cell(row=fila, column=3, value=detalle)
                c2.font = fuente(color='334155', size=9)
                c2.alignment = izquierda(True)
                c2.border = borde_fino()
                ws.cell(row=fila, column=4).border = borde_fino()
                ws.row_dimensions[fila].height = _altura((detalle, 68), (etiqueta, 18))
                fila += 1
            fila += 1  # separador entre secciones

        # ── Leyenda: "CÓMO LEER ESTE REPORTE" ────────────────────────────────
        ws.merge_cells(f'A{fila}:D{fila}')
        c = ws.cell(row=fila, column=1, value='📖  CÓMO LEER ESTE REPORTE')
        c.font = fuente(bold=True, color=BLANCO, size=11)
        c.fill = fill(AZUL_OSCURO)
        c.alignment = centrado()
        c.border = borde_fino()
        ws.row_dimensions[fila].height = 22
        fila += 1
        leyenda = [
            ('Turno normal', 'DBEAFE', '1D4ED8', 'Trabaja su jornada de siempre (AM o PM).'),
            ('Dobla',        'EDE9FE', '5B21B6', 'Trabaja AM y PM para cubrir a un compañero que descansa.'),
            ('Cambio',       'FEF9C3', '92400E', 'Trabaja por un cambio de turno aprobado.'),
            ('💥 Hoy paga doblada', 'DCFCE7', '166534', 'Hoy cumple una doblada que debía de otro día.'),
            ('🔁 Debe doblada', 'EDE9FE', '5B21B6', 'No cumplió un día de doblada; está pendiente de pagarlo.'),
            ('⚠ Restricción', 'FEF3C7', '92400E', 'Tiene una restricción médica/operativa activa (ver recomendación).'),
            ('⛔ Sanción',     'FEE2E2', '991B1B', 'Está sancionado/a; no puede hacer solicitudes mientras dure.'),
            ('🕐 Permiso',     'FEFCE8', '854D0E', 'Tiene permiso especial ese día (horas indicadas).'),
            ('⚠ Sin planificar', 'FEE2E2', '991B1B',
             'El día es finde o festivo y todavía no tiene publicada la alternancia: '
             'las columnas vacías NO significan que descansen todos.'),
            ('Con quién', 'F1F5F9', '334155',
             'Compañero del acuerdo que puso a esa persona en la jornada de hoy. Al lado: '
             'el tipo de solicitud, qué papel jugó y "Fecha relacionada", que es el OTRO '
             'día del trato (si hoy cede, el día en que se lo devuelven, y al revés).'),
            ('Cambios y permisos', 'F1F5F9', '334155',
             'Hoja aparte con UNA fila por movimiento del día: quién lo pidió, con quién, '
             'qué pasa hoy y cuándo se aprobó. Es el resumen para planear la semana.'),
        ]
        for etiqueta, bg, fg, desc_txt in leyenda:
            cb = ws.cell(row=fila, column=1, value=etiqueta)
            cb.font = fuente(bold=True, color=fg, size=9)
            cb.fill = fill(bg)
            cb.alignment = centrado(True)
            cb.border = borde_fino()
            ws.merge_cells(f'B{fila}:D{fila}')
            cd = ws.cell(row=fila, column=2, value=desc_txt)
            cd.font = fuente(color='334155', size=9)
            cd.alignment = izquierda(True)
            cd.border = borde_fino()
            for cc in (3, 4):
                ws.cell(row=fila, column=cc).border = borde_fino()
            ws.row_dimensions[fila].height = _altura((desc_txt, 68))
            fila += 1

        # ════════════════════════════════════════════════════════════════════
        # HOJA 2 — TRABAJAN AM
        # ════════════════════════════════════════════════════════════════════
        ws_am = wb.create_sheet('Trabajan AM')
        self._hoja_trabajan(ws_am, 'TRABAJAN AM ☀️', am, fecha_legible,
                            AZUL_MEDIO, AZUL_CLARO, AZUL_OSC=AZUL_OSCURO,
                            fill=fill, fuente=fuente, borde=borde_fino,
                            centrado=centrado, izquierda=izquierda, dia_info=dia_info)

        # ════════════════════════════════════════════════════════════════════
        # HOJA 3 — TRABAJAN PM
        # ════════════════════════════════════════════════════════════════════
        ws_pm = wb.create_sheet('Trabajan PM')
        self._hoja_trabajan(ws_pm, 'TRABAJAN PM 🌙', pm, fecha_legible,
                            NARANJA_MED, NARANJA_CL, AZUL_OSC=NARANJA_OSC,
                            fill=fill, fuente=fuente, borde=borde_fino,
                            centrado=centrado, izquierda=izquierda, dia_info=dia_info)

        # ════════════════════════════════════════════════════════════════════
        # HOJA 4 — DESCANSAN
        # ════════════════════════════════════════════════════════════════════
        ws_desc = wb.create_sheet('Descansan')
        self._hoja_descansan(ws_desc, desc, fecha_legible,
                             fill=fill, fuente=fuente, borde=borde_fino,
                             centrado=centrado, izquierda=izquierda,
                             GRIS_OSC=GRIS_OSC, GRIS_MED=GRIS_MED,
                             GRIS_CLARO=GRIS_CLARO, VERDE_CL=VERDE_CL,
                             MORADO_CL=MORADO_CL)

        # ════════════════════════════════════════════════════════════════════
        # HOJA 5 — CAMBIOS Y PERMISOS DEL DÍA
        # ════════════════════════════════════════════════════════════════════
        ws_cambios = wb.create_sheet('Cambios y permisos')
        self._hoja_cambios(ws_cambios, trabajando, descansando, fecha_legible,
                           fill=fill, fuente=fuente, borde=borde_fino,
                           centrado=centrado, izquierda=izquierda,
                           VERDE_OSC='065F46', VERDE_MED='0D9488')

        # ════════════════════════════════════════════════════════════════════
        # HOJA 6 — DEUDA PENDIENTE AL CORTE
        # ════════════════════════════════════════════════════════════════════
        ws_deuda = wb.create_sheet('Deuda pendiente')
        self._hoja_deuda(ws_deuda, filas_deuda or [], resumen_deuda or {},
                         fecha, fecha_legible,
                         fill=fill, fuente=fuente, borde=borde_fino,
                         centrado=centrado, izquierda=izquierda,
                         ROJO_OSC=ROJO_OSC, ROJO_MED=ROJO_MED)

        # ── Respuesta HTTP ───────────────────────────────────────────────────
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        nombre_archivo = f'reporte_operacion_{fecha}.xlsx'
        response = HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}"'
        return response

    def _encabezado_hoja(self, ws, titulo, fecha_legible, color_fondo, color_texto,
                          fill, fuente, centrado, cols):
        from openpyxl.utils import get_column_letter
        ws.sheet_view.showGridLines = False
        ultima = get_column_letter(len(cols))
        ws.merge_cells(f'A1:{ultima}1')
        c = ws['A1']
        c.value = titulo
        c.font = fuente(bold=True, color='FFFFFF', size=13)
        c.fill = fill(color_fondo)
        c.alignment = centrado()
        ws.row_dimensions[1].height = 28
        ws.merge_cells(f'A2:{ultima}2')
        c = ws['A2']
        c.value = fecha_legible.capitalize()
        c.font = fuente(italic=True, color='FFFFFF', size=10)
        c.fill = fill(color_texto)
        c.alignment = centrado()
        ws.row_dimensions[2].height = 18
        ws.row_dimensions[3].height = 6

    def _fila_header(self, ws, fila, cols, color_fondo, color_texto,
                     fill, fuente, centrado, borde):
        from openpyxl.utils import get_column_letter
        for col, (ancho, label) in enumerate(cols, 1):
            c = ws.cell(row=fila, column=col, value=label)
            c.font = fuente(bold=True, color='FFFFFF', size=10)
            c.fill = fill(color_fondo)
            c.alignment = centrado(True)
            c.border = borde()
            ws.row_dimensions[fila].height = 26
            ws.column_dimensions[get_column_letter(col)].width = ancho

    def _hoja_trabajan(self, ws, titulo, empleados, fecha_legible,
                       color_enc, color_sub, AZUL_OSC,
                       fill, fuente, borde, centrado, izquierda, dia_info):
        from openpyxl.utils import get_column_letter
        COLS = [(5, '#'), (20, 'Nombre'), (16, 'Apellido'), (12, 'Jornada hoy'),
                (54, '¿Por qué trabaja hoy?'),
                (24, 'Con quién'), (20, 'Tipo de cambio'), (17, 'Su papel'),
                (16, 'Fecha relacionada'), (18, 'Aprobado el'),
                (34, 'Permiso'), (34, 'Restricción'), (30, 'Sanción'),
                (36, 'Doblada pendiente')]
        # Índices POR NOMBRE, no a mano: el pintado de abajo iba con `col == 6/7/8/9`
        # hardcodeado, así que insertar una columna en medio descolocaba en silencio todos
        # los colores. Derivándolos de COLS, añadir o mover una columna es cambiar una línea.
        IDX = {label: n for n, (_, label) in enumerate(COLS, 1)}
        self._encabezado_hoja(ws, titulo, fecha_legible, AZUL_OSC, color_enc,
                               fill, fuente, centrado, COLS)
        self._fila_header(ws, 4, COLS, color_enc, 'FFFFFF',
                          fill, fuente, centrado, borde)
        TIPO_BG = {'oficial': 'DBEAFE', 'cambio': 'FEF9C3', 'doblada': 'EDE9FE'}
        TIPO_FG = {'oficial': '1D4ED8', 'cambio': '92400E', 'doblada': '5B21B6'}
        for i, emp in enumerate(empleados, 5):
            bg = 'FFFFFF' if i % 2 == 1 else 'F8FAFC'
            tipo = emp.get('tipo', 'oficial')
            deuda = emp.get('deuda_reprogramacion') or {}
            frase = _frase_trabaja(emp, dia_info)
            acu = _celdas_acuerdo(emp)
            hay_acuerdo = bool(emp.get('acuerdo'))
            t_perm, t_rest, t_sanc, t_deuda = (_txt_permiso(emp), _txt_restriccion(emp),
                                               _txt_sancion(emp), _txt_deuda(emp))
            vals = [i - 4, emp['nombre'], emp['apellido'], emp.get('jornada_dia', ''),
                    frase, *acu,
                    t_perm or '—', t_rest or '—', t_sanc or '—', t_deuda or '—']
            for col, v in enumerate(vals, 1):
                c = ws.cell(row=i, column=col, value=v)
                c.border = borde()
                c.alignment = izquierda(True) if col > 1 else centrado()
                c.font = fuente(size=10)
                if col == IDX['¿Por qué trabaja hoy?']:  # frase: color según tipo
                    c.fill = fill(TIPO_BG.get(tipo, bg))
                    c.font = fuente(bold=True, color=TIPO_FG.get(tipo, '000000'), size=10)
                elif col == IDX['Con quién']:
                    c.fill = fill(TIPO_BG.get(tipo, bg) if hay_acuerdo else bg)
                    if hay_acuerdo:
                        c.font = fuente(bold=True, color=TIPO_FG.get(tipo, '000000'), size=10)
                elif col == IDX['Tipo de cambio']:
                    c.fill = fill(TIPO_BG.get(tipo, bg) if hay_acuerdo else bg)
                elif col == IDX['Permiso']:
                    c.fill = fill('FEFCE8' if t_perm else bg)
                elif col == IDX['Restricción']:
                    c.fill = fill('FEF3C7' if t_rest else bg)
                elif col == IDX['Sanción']:
                    c.fill = fill('FEE2E2' if t_sanc else bg)
                elif col == IDX['Doblada pendiente']:
                    if deuda.get('paga_hoy'):
                        c.fill = fill('DCFCE7'); c.font = fuente(bold=True, size=10)
                    else:
                        c.fill = fill('EDE9FE' if t_deuda else bg)
                else:
                    c.fill = fill(bg)
            ws.row_dimensions[i].height = _altura(
                (frase, 54), (acu[0], 24), (t_perm, 34), (t_rest, 34),
                (t_sanc, 30), (t_deuda, 36))
        if not empleados:
            ws.merge_cells(f'A5:{get_column_letter(len(COLS))}5')
            c = ws['A5']
            c.value = 'No hay exploradores en esta jornada'
            c.font = fuente(italic=True, color='94A3B8', size=10)
            c.alignment = centrado()

    def _hoja_descansan(self, ws, empleados, fecha_legible,
                        fill, fuente, borde, centrado, izquierda,
                        GRIS_OSC, GRIS_MED, GRIS_CLARO, VERDE_CL, MORADO_CL):
        from openpyxl.utils import get_column_letter
        COLS = [(5, '#'), (20, 'Nombre'), (16, 'Apellido'), (11, 'Jornada base'),
                (56, '¿Por qué descansa hoy?'),
                (24, 'Con quién'), (20, 'Tipo de cambio'), (17, 'Su papel'),
                (16, 'Fecha relacionada'), (18, 'Aprobado el'),
                (34, 'Permiso'), (34, 'Restricción'), (30, 'Sanción'),
                (36, 'Doblada pendiente')]
        IDX = {label: n for n, (_, label) in enumerate(COLS, 1)}
        self._encabezado_hoja(ws, 'DESCANSAN 🛌', fecha_legible, GRIS_OSC, GRIS_MED,
                               fill, fuente, centrado, COLS)
        self._fila_header(ws, 4, COLS, GRIS_MED, 'FFFFFF',
                          fill, fuente, centrado, borde)
        MOTIVO_BG = {
            'cedió su jornada':          MORADO_CL,
            'paga doblada':              MORADO_CL,
            'doblada permanente':        MORADO_CL,
            'cambio de día de descanso': VERDE_CL,
            'descanso de temporada':     'FEF9C3',
            'lunes de mantenimiento':    GRIS_CLARO,
            'descanso de fin de semana': 'F0FDF4',
            'festivo':                   'FEF2F2',
            'sin alternancia publicada': 'FEE2E2',
        }
        for i, emp in enumerate(empleados, 5):
            bg = 'FFFFFF' if i % 2 == 1 else 'F8FAFC'
            motivo = (emp.get('motivo') or '').lower()
            deuda = emp.get('deuda_reprogramacion') or {}
            frase = _frase_descansa(emp)
            acu = _celdas_acuerdo(emp)
            hay_acuerdo = bool(emp.get('acuerdo'))
            t_perm, t_rest, t_sanc, t_deuda = (_txt_permiso(emp), _txt_restriccion(emp),
                                               _txt_sancion(emp), _txt_deuda(emp))
            vals = [i - 4, emp['nombre'], emp['apellido'], emp.get('jornada_base', ''),
                    frase, *acu,
                    t_perm or '—', t_rest or '—', t_sanc or '—', t_deuda or '—']
            motivo_bg = MOTIVO_BG.get(motivo, bg)
            for col, v in enumerate(vals, 1):
                c = ws.cell(row=i, column=col, value=v)
                c.border = borde()
                c.alignment = izquierda(True) if col > 1 else centrado()
                c.font = fuente(size=10)
                if col == IDX['¿Por qué descansa hoy?']:  # frase: color según motivo
                    c.fill = fill(motivo_bg)
                    c.font = fuente(bold=True, color=GRIS_OSC, size=10)
                elif col == IDX['Con quién']:
                    c.fill = fill(motivo_bg if hay_acuerdo else bg)
                    if hay_acuerdo:
                        c.font = fuente(bold=True, color=GRIS_OSC, size=10)
                elif col == IDX['Tipo de cambio']:
                    c.fill = fill(motivo_bg if hay_acuerdo else bg)
                elif col == IDX['Permiso']:
                    c.fill = fill('FEFCE8' if t_perm else bg)
                elif col == IDX['Restricción']:
                    c.fill = fill('FEF3C7' if t_rest else bg)
                elif col == IDX['Sanción']:
                    c.fill = fill('FEE2E2' if t_sanc else bg)
                elif col == IDX['Doblada pendiente']:
                    if deuda.get('paga_hoy'):
                        c.fill = fill('DCFCE7'); c.font = fuente(bold=True, size=10)
                    else:
                        c.fill = fill('EDE9FE' if t_deuda else bg)
                else:
                    c.fill = fill(bg)
            ws.row_dimensions[i].height = _altura(
                (frase, 56), (acu[0], 24), (t_perm, 34), (t_rest, 34),
                (t_sanc, 30), (t_deuda, 36))
        if not empleados:
            ws.merge_cells(f'A5:{get_column_letter(len(COLS))}5')
            c = ws['A5']
            c.value = 'Todos los exploradores trabajan este día'
            c.font = fuente(italic=True, color='94A3B8', size=10)
            c.alignment = centrado()

    def _hoja_cambios(self, ws, trabajando, descansando, fecha_legible,
                      fill, fuente, borde, centrado, izquierda, VERDE_OSC, VERDE_MED):
        """UNA fila por MOVIMIENTO del día: qué se movió, entre quiénes y cuándo se aprobó.

        Las otras hojas contestan "¿qué hace fulano hoy?". Esta contesta la pregunta del
        supervisor que planea la semana: "¿qué se movió y con quién queda pendiente la
        devolución?".

        Un acuerdo aparece en las DOS personas (quien trabaja y quien descansa son las dos
        caras del mismo trato), así que se DEDUPLICA por número de solicitud: es un
        movimiento, no dos. La columna "Qué pasa hoy" junta los dos lados en una frase.
        """
        from openpyxl.utils import get_column_letter
        COLS = [(5, '#'), (20, 'Tipo'), (22, 'Quién pidió'), (22, 'Con quién'),
                (52, 'Qué pasa hoy'), (16, 'Fecha relacionada'),
                (18, 'Solicitado el'), (18, 'Aprobado el'), (22, 'Aprobado por'),
                (13, 'N° solicitud')]
        self._encabezado_hoja(ws, 'CAMBIOS Y PERMISOS DEL DÍA 🔄', fecha_legible,
                              VERDE_OSC, VERDE_MED, fill, fuente, centrado, COLS)
        self._fila_header(ws, 4, COLS, VERDE_MED, 'FFFFFF',
                          fill, fuente, centrado, borde)

        def _persona(emp):
            return f'{emp["nombre"]} {emp["apellido"]}'.strip()

        movimientos = {}
        for emp, verbo in ([(e, 'trabaja') for e in trabajando]
                           + [(e, 'descansa') for e in descansando]):
            a = emp.get('acuerdo') or {}
            if a.get('solicitud_id'):
                clave = ('sol', a['solicitud_id'])
                mov = movimientos.get(clave)
                if mov is None:
                    # `rol` desconocido (dato antiguo) se trata como solicitante: es el
                    # caso mayoritario y deja la fila legible en vez de invertida.
                    es_sol = a.get('rol') != 'receptor'
                    yo, otro = _persona(emp), a.get('companero_nombre') or '—'
                    mov = movimientos[clave] = {
                        'orden': (0, a.get('tipo') or '', yo),
                        'vals': [a.get('tipo') or '—',
                                 yo if es_sol else otro,
                                 otro if es_sol else yo,
                                 [],
                                 a.get('fecha_relacionada') or '—',
                                 a.get('fecha_solicitud') or '—',
                                 a.get('fecha_resolucion') or '—',
                                 '—',
                                 a['solicitud_id']],
                    }
                jd = emp.get('jornada_dia') if verbo == 'trabaja' else None
                mov['vals'][3].append(f'{_persona(emp)} {verbo}'
                                      + (f' {jd}' if jd else ''))
            p = emp.get('permiso') or {}
            if p.get('id'):
                horas = p.get('horas', 0) or 0
                esp = f' — {p["especificacion"]}' if p.get('especificacion') else ''
                cubre = (p.get('cubre') or {}).get('nombre') or '—'
                movimientos[('permiso', p['id'])] = {
                    'orden': (1, 'PERMISO', _persona(emp)),
                    'vals': ['PERMISO', _persona(emp), cubre,
                             [f'{_persona(emp)}: permiso de {horas:g}h '
                              f'({p.get("tipo", "")}{esp}) — {p.get("estado", "")}'],
                             '—',
                             p.get('fecha_solicitud') or '—',
                             p.get('fecha_aprobacion') or '—',
                             (p.get('aprobado_por') or {}).get('nombre') or '—',
                             p['id']],
                }

        filas = sorted(movimientos.values(), key=lambda m: m['orden'])
        for i, mov in enumerate(filas, 5):
            bg = 'FFFFFF' if i % 2 == 1 else 'F8FAFC'
            vals = list(mov['vals'])
            vals[3] = ' · '.join(vals[3]) or '—'
            es_permiso = vals[0] == 'PERMISO'
            for col, v in enumerate([i - 4, *vals], 1):
                c = ws.cell(row=i, column=col, value=v)
                c.border = borde()
                c.alignment = izquierda(True) if col > 1 else centrado()
                c.font = fuente(size=10)
                if col == 2:
                    c.fill = fill('FEFCE8' if es_permiso else 'CCFBF1')
                    c.font = fuente(bold=True, color='854D0E' if es_permiso else VERDE_OSC,
                                    size=10)
                else:
                    c.fill = fill(bg)
            ws.row_dimensions[i].height = _altura((vals[3], 52), (vals[1], 22), (vals[2], 22))

        if not filas:
            ws.merge_cells(f'A5:{get_column_letter(len(COLS))}5')
            c = ws['A5']
            c.value = 'No hay cambios ni permisos aprobados que afecten este día'
            c.font = fuente(italic=True, color='94A3B8', size=10)
            c.alignment = centrado()

        # Nota al pie: por qué "Aprobado por" solo se llena en los permisos.
        fila = 5 + max(len(filas), 1) + 1
        ws.merge_cells(f'A{fila}:{get_column_letter(len(COLS))}{fila}')
        c = ws.cell(row=fila, column=1, value=(
            'ℹ  "Aprobado por" solo aparece en los permisos: de las solicitudes de cambio se '
            'guarda CUÁNDO quedaron aprobadas, no qué supervisor las aprobó. '
            'El "N° solicitud" sirve para buscarla en el módulo de solicitudes.'))
        c.font = fuente(italic=True, color='64748B', size=9)
        c.alignment = izquierda(True)
        ws.row_dimensions[fila].height = _altura((c.value, 120))

    def _hoja_deuda(self, ws, filas, resumen, fecha, fecha_legible,
                    fill, fuente, borde, centrado, izquierda, ROJO_OSC, ROJO_MED):
        """
        Deuda todavía sin pagar generada del día 1 de ese mes hasta la fecha elegida.

        Es la misma lectura que la pantalla de morosos con fecha de corte: solo lo ya
        devengado (de un permiso permanente entran únicamente las ocurrencias que ya
        pasaron) y sin arrastrar meses anteriores, que a estas alturas o están pagados
        o los consumió una sanción.

        Con fecha futura la cifra deja de ser exigible y se rotula PROYECTADA: aquí el
        supervisor llega desde un calendario donde clicar un día que aún no ha llegado es
        lo normal, y "debe 3 h" de un día que no ha pasado se lee mal.
        """
        from openpyxl.utils import get_column_letter

        proyectada = bool(resumen.get('proyectada'))
        periodo = f'del 1 al {fecha.day} de {MESES_ES[fecha.month]} de {fecha.year}'

        COLS = [(5, '#'), (20, 'Nombre'), (16, 'Apellido'), (24, 'Supervisor'),
                (12, 'Debe (h)'), (12, 'Debe (min)'), (18, 'Dobladas (min)'),
                (18, 'Permisos (min)'), (12, 'Ocasiones'), (14, 'Desde')]
        titulo = (f'DEUDA {"PROYECTADA ⏳" if proyectada else "PENDIENTE 💸"}  ·  {periodo}')
        self._encabezado_hoja(ws, titulo, fecha_legible, ROJO_OSC, ROJO_MED,
                              fill, fuente, centrado, COLS)
        self._fila_header(ws, 4, COLS, ROJO_MED, 'FFFFFF',
                          fill, fuente, centrado, borde)

        ultima = get_column_letter(len(COLS))
        for i, f in enumerate(filas, 5):
            bg = 'FFFFFF' if i % 2 == 1 else 'F8FAFC'
            explorador = f['explorador']
            sup = explorador.supervisor
            vals = [i - 4, explorador.nombre, explorador.apellido,
                    f'{sup.nombre} {sup.apellido}' if sup else '—',
                    f['horas'], f['minutos'],
                    f['minutos_dobladas'] or '—', f['minutos_permisos'] or '—',
                    f['dias'], _fmt_fecha(f['deuda_mas_antigua'])]
            for col, v in enumerate(vals, 1):
                c = ws.cell(row=i, column=col, value=v)
                c.border = borde()
                c.font = fuente(size=10)
                c.alignment = izquierda(True) if col in (2, 3, 4) else centrado()
                if col == 5:  # las horas son la cifra que se paga: se destacan
                    c.fill = fill('FEE2E2')
                    c.font = fuente(bold=True, color=ROJO_OSC, size=10)
                else:
                    c.fill = fill(bg)

        if not filas:
            ws.merge_cells(f'A5:{ultima}5')
            c = ws['A5']
            c.value = f'Nadie debe nada {periodo} ✅'
            c.font = fuente(italic=True, color='94A3B8', size=10)
            c.alignment = centrado()
            return

        total_min = resumen.get('minutos') or sum(f['minutos'] for f in filas)
        fila_total = 5 + len(filas)
        ws.merge_cells(f'A{fila_total}:D{fila_total}')
        c = ws[f'A{fila_total}']
        c.value = f'TOTAL · {len(filas)} explorador(es) deben'
        c.font = fuente(bold=True, color='FFFFFF', size=10)
        c.fill = fill(ROJO_OSC)
        c.alignment = centrado()
        for col, v in enumerate([round(total_min / 60, 2), total_min,
                                 sum(f['minutos_dobladas'] for f in filas),
                                 sum(f['minutos_permisos'] for f in filas),
                                 sum(f['dias'] for f in filas), ''], 5):
            c = ws.cell(row=fila_total, column=col, value=v)
            c.font = fuente(bold=True, color='FFFFFF', size=10)
            c.fill = fill(ROJO_OSC)
            c.alignment = centrado()
            c.border = borde()

        nota_txt = ('Solo cuenta lo ya devengado a esa fecha: de un permiso permanente entran '
                    'únicamente las ocurrencias que ya han pasado. La deuda de meses anteriores '
                    'no aparece — a estas alturas o está pagada o la consumió una sanción.')
        if proyectada:
            nota_txt = ('⏳ PROYECCIÓN: la fecha elegida aún no ha llegado, así que esta cifra '
                        'incluye días que todavía no han vencido. No es deuda exigible hoy. '
                        + nota_txt)
        nota = fila_total + 2
        ws.merge_cells(f'A{nota}:{ultima}{nota}')
        c = ws[f'A{nota}']
        c.value = nota_txt
        c.font = fuente(italic=True, color=('991B1B' if proyectada else '64748B'), size=9)
        c.alignment = izquierda(True)
        ws.row_dimensions[nota].height = 30
