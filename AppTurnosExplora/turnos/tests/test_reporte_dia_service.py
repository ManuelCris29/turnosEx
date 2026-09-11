"""
Tests del enriquecimiento del reporte del día: restricción, sanción y deuda de
doblada por reprogramación. Verifica que las claves nuevas aparezcan (o queden
None) según los datos, y un smoke test de la exportación a Excel.
"""
import io
from datetime import date, timedelta
from unittest.mock import patch

from django.contrib.auth.models import User
from django.test import RequestFactory, TestCase

from empleados.models import Empleado, Jornada, RestriccionEmpleado, SancionEmpleado
from solicitudes.models import ReprogramacionDiaDoblada, SolicitudCambio, TipoSolicitudCambio
from turnos.models import AsignarJornadaExplorador, Sala
from turnos.services.reporte_dia_service import ReporteDiaService

FECHA = date(2026, 3, 4)  # miércoles, no festivo


def _buscar(data, emp_id):
    """Encuentra el dict del empleado en trabajando o descansando."""
    for grupo in ('trabajando', 'descansando'):
        for e in data[grupo]:
            if e['id'] == emp_id:
                return e
    return None


class ReporteDiaEnriquecidoTest(TestCase):

    def setUp(self):
        from django.core.cache import cache
        cache.clear()
        self.am = Jornada.objects.create(nombre='AM', hora_inicio='06:00:00', hora_fin='14:00:00')
        self.pm = Jornada.objects.create(nombre='PM', hora_inicio='14:00:00', hora_fin='22:00:00')
        self.sala = Sala.objects.create(nombre='Sala Rep', activo=True)
        self.tipo_dob = TipoSolicitudCambio.objects.create(nombre='DOBLADA', codigo_estrategia='DOBLADA', activo=True)

        def _emp(username, ced, jor, staff=False):
            u = User.objects.create_user(username=username, password='x', is_staff=staff)
            e = Empleado.objects.create(user=u, nombre=username, apellido='X', cedula=ced, activo=True)
            AsignarJornadaExplorador.objects.create(explorador=e, jornada=jor, fecha_inicio=date(2025, 1, 1))
            return e

        self.a = _emp('rep_a', '8001', self.am)
        self.b = _emp('rep_b', '8002', self.pm)
        self.sup = Empleado.objects.create(
            user=User.objects.create_user(username='rep_sup', password='x', is_staff=True),
            nombre='Sup', apellido='X', cedula='8003', activo=True,
        )

    # ── 1. Claves nuevas presentes y None sin datos ──────────────────────────
    def test_claves_nuevas_none_sin_datos(self):
        data = ReporteDiaService.reporte(FECHA)
        ea = _buscar(data, self.a.id)
        for k in ('restriccion', 'sancion', 'deuda_reprogramacion'):
            self.assertIn(k, ea)
            self.assertIsNone(ea[k])

    # ── 2. Restricciones ─────────────────────────────────────────────────────
    def test_restriccion_activa_y_vencida(self):
        RestriccionEmpleado.objects.create(
            empleado=self.a, fecha_inicio=FECHA - timedelta(5), fecha_fin=FECHA + timedelta(5),
            tipo_restriccion='Física', recomendacion='Sin cargas pesadas',
        )
        # Vencida para B (no debe aparecer)
        RestriccionEmpleado.objects.create(
            empleado=self.b, fecha_inicio=FECHA - timedelta(20), fecha_fin=FECHA - timedelta(1),
            tipo_restriccion='Vieja', recomendacion='ya pasó',
        )
        data = ReporteDiaService.reporte(FECHA)
        self.assertEqual(_buscar(data, self.a.id)['restriccion']['tipo'], 'Física')
        self.assertEqual(_buscar(data, self.a.id)['restriccion']['recomendacion'], 'Sin cargas pesadas')
        self.assertIsNone(_buscar(data, self.b.id)['restriccion'])

    def test_restriccion_indefinida(self):
        RestriccionEmpleado.objects.create(
            empleado=self.a, fecha_inicio=FECHA - timedelta(1), fecha_fin=None,
            tipo_restriccion='Perm', recomendacion='indefinida',
        )
        r = _buscar(ReporteDiaService.reporte(FECHA), self.a.id)['restriccion']
        self.assertIsNotNone(r)
        self.assertIsNone(r['fecha_fin'])

    # ── 3. Sanciones ─────────────────────────────────────────────────────────
    def test_sancion_activa_indefinida_y_futura(self):
        SancionEmpleado.objects.create(
            explorador=self.a, supervisor=self.sup, fecha_inicio=FECHA - timedelta(2),
            fecha_fin=None, motivo='Deuda de horas',
        )
        # Futura para B (no debe aparecer)
        SancionEmpleado.objects.create(
            explorador=self.b, supervisor=self.sup, fecha_inicio=FECHA + timedelta(5),
            fecha_fin=None, motivo='futura',
        )
        data = ReporteDiaService.reporte(FECHA)
        self.assertEqual(_buscar(data, self.a.id)['sancion']['motivo'], 'Deuda de horas')
        self.assertIsNone(_buscar(data, self.b.id)['sancion'])

    # ── 4-6. Deuda de doblada por reprogramación ─────────────────────────────
    def _reprog(self, **kwargs):
        origen = SolicitudCambio.objects.create(
            explorador_solicitante=self.a, explorador_receptor=self.b,
            tipo_cambio=self.tipo_dob, fecha_cambio_turno=FECHA - timedelta(30), estado='aprobada',
        )
        defaults = dict(doblada_origen=origen, explorador=self.a,
                        fecha_original=FECHA - timedelta(10), jornada_debida='PM', motivo='incapacidad')
        defaults.update(kwargs)
        return ReprogramacionDiaDoblada.objects.create(**defaults)

    def test_deuda_pendiente_sin_fecha(self):
        self._reprog(estado='pendiente', fecha_reprogramada=None)
        d = _buscar(ReporteDiaService.reporte(FECHA), self.a.id)['deuda_reprogramacion']
        self.assertEqual(d['estado'], 'pendiente')
        self.assertIsNone(d['fecha_reprogramada'])
        self.assertFalse(d['paga_hoy'])

    def test_deuda_paga_hoy(self):
        self._reprog(estado='pagada', fecha_reprogramada=FECHA)
        d = _buscar(ReporteDiaService.reporte(FECHA), self.a.id)['deuda_reprogramacion']
        self.assertTrue(d['paga_hoy'])

    def test_deuda_pagada_en_pasado_no_aparece(self):
        self._reprog(estado='pagada', fecha_reprogramada=FECHA - timedelta(2))
        self.assertIsNone(_buscar(ReporteDiaService.reporte(FECHA), self.a.id)['deuda_reprogramacion'])

    # ── 7. Smoke test del Excel ──────────────────────────────────────────────
    def test_excel_se_genera_con_hojas_y_headers(self):
        import openpyxl

        from turnos.api.views.reportes import ReporteDiaExcelView
        req = RequestFactory().get(f'/x?fecha={FECHA.isoformat()}')
        req.user = self.sup.user
        resp = ReporteDiaExcelView().get(req)
        self.assertEqual(resp.status_code, 200)
        self.assertIn('spreadsheetml', resp['Content-Type'])
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertEqual(wb.sheetnames,
                         ['Resumen', 'Trabajan AM', 'Trabajan PM', 'Descansan',
                          'Cambios y permisos', 'Deuda pendiente'])
        deuda = [wb['Deuda pendiente'].cell(row=4, column=c).value for c in range(1, 11)]
        self.assertIn('Debe (h)', deuda)
        self.assertIn('Supervisor', deuda)
        self.assertIn('Dobladas (min)', deuda)
        self.assertIn('Permisos (min)', deuda)
        headers = [wb['Trabajan AM'].cell(row=4, column=c).value for c in range(1, 15)]
        self.assertIn('¿Por qué trabaja hoy?', headers)
        self.assertIn('Restricción', headers)
        self.assertIn('Doblada pendiente', headers)
        # El «con quién» del acuerdo: columnas propias, no enterrado en la frase.
        for h in ('Con quién', 'Tipo de cambio', 'Su papel', 'Fecha relacionada',
                  'Aprobado el'):
            self.assertIn(h, headers)
        self.assertIn('Con quién',
                      [wb['Descansan'].cell(row=4, column=c).value for c in range(1, 15)])

    def test_excel_hoja_deuda_lista_lo_pendiente_del_mes_hasta_la_fecha(self):
        """La hoja replica /empleados/sanciones/morosos/?corte=: del día 1 al día elegido."""
        import openpyxl

        from solicitudes.models import DeudaCorporativa
        from turnos.api.views.reportes import ReporteDiaExcelView

        DeudaCorporativa.objects.create(explorador=self.a, minutos=30,
                                        fecha_doblada=FECHA - timedelta(1), estado='activa')
        # Posterior al corte: todavía no se debe, no debe aparecer.
        DeudaCorporativa.objects.create(explorador=self.b, minutos=30,
                                        fecha_doblada=FECHA + timedelta(1), estado='activa')

        req = RequestFactory().get(f'/x?fecha={FECHA.isoformat()}')
        req.user = self.sup.user
        wb = openpyxl.load_workbook(io.BytesIO(ReporteDiaExcelView().get(req).content))
        ws = wb['Deuda pendiente']

        nombres = [ws.cell(row=r, column=2).value for r in range(5, ws.max_row + 1)]
        self.assertIn(self.a.nombre, nombres)
        self.assertNotIn(self.b.nombre, nombres)
        self.assertEqual(ws.cell(row=5, column=6).value, 30)  # minutos

    def test_excel_hoja_deuda_avisa_cuando_la_fecha_es_futura(self):
        """Una fecha que aún no ha llegado no es deuda exigible: se rotula PROYECTADA."""
        import openpyxl

        from turnos.api.views.reportes import ReporteDiaExcelView

        req = RequestFactory().get('/x?fecha=' + FECHA.isoformat())
        req.user = self.sup.user
        with patch('django.utils.timezone.localdate', return_value=FECHA - timedelta(5)):
            wb = openpyxl.load_workbook(io.BytesIO(ReporteDiaExcelView().get(req).content))
        self.assertIn('PROYECTADA', wb['Deuda pendiente']['A1'].value)

        with patch('django.utils.timezone.localdate', return_value=FECHA + timedelta(5)):
            wb = openpyxl.load_workbook(io.BytesIO(ReporteDiaExcelView().get(req).content))
        self.assertIn('PENDIENTE', wb['Deuda pendiente']['A1'].value)

    def test_excel_hoja_deuda_trae_el_supervisor(self):
        import openpyxl

        from solicitudes.models import DeudaCorporativa
        from turnos.api.views.reportes import ReporteDiaExcelView

        self.a.supervisor = self.sup
        self.a.save(update_fields=['supervisor'])
        DeudaCorporativa.objects.create(explorador=self.a, minutos=30,
                                        fecha_doblada=FECHA, estado='activa')

        req = RequestFactory().get('/x?fecha=' + FECHA.isoformat())
        req.user = self.sup.user
        wb = openpyxl.load_workbook(io.BytesIO(ReporteDiaExcelView().get(req).content))
        self.assertEqual(wb['Deuda pendiente'].cell(row=5, column=4).value,
                         f'{self.sup.nombre} {self.sup.apellido}')

    def test_json_trae_el_agregado_de_deuda(self):
        """La pantalla necesita el titular; el detalle nominal sigue solo en morosos."""
        import json

        from solicitudes.models import DeudaCorporativa
        from turnos.api.views.reportes import ReporteDiaView

        DeudaCorporativa.objects.create(explorador=self.a, minutos=90,
                                        fecha_doblada=FECHA, estado='activa')
        req = RequestFactory().get('/x?fecha=' + FECHA.isoformat())
        req.user = self.sup.user
        data = json.loads(ReporteDiaView().get(req).content)

        self.assertEqual(data['deuda_corte'],
                         {'exploradores': 1, 'minutos': 90, 'horas': 1.5,
                          'proyectada': FECHA > date.today()})
        # Solo el agregado: ningún nombre viaja al JSON de la pantalla.
        self.assertNotIn(self.a.nombre, json.dumps(data['deuda_corte']))

    # ── 7b. Tipo de día del mes (pintado del calendario) ─────────────────────
    def test_dias_del_mes_marca_festivo_mantenimiento_y_sin_planificar(self):
        from turnos.models import AsignacionEspecialManual, DiaEspecial

        festivo = date(2026, 3, 23)   # lunes
        mant = date(2026, 3, 30)      # lunes
        DiaEspecial.objects.create(fecha=festivo, tipo='festivo', activo=True)
        DiaEspecial.objects.create(fecha=mant, tipo='mantenimiento', activo=True)
        sabado = date(2026, 3, 7)
        AsignacionEspecialManual.objects.create(fecha=sabado, jornada_trabaja=self.am, activo=True)

        dias = ReporteDiaService.dias_del_mes(2026, 3)

        self.assertTrue(dias[festivo.isoformat()]['es_festivo'])
        self.assertTrue(dias[festivo.isoformat()]['sin_planificar'])  # festivo entre semana sin alternancia
        self.assertTrue(dias[mant.isoformat()]['es_mantenimiento'])
        self.assertFalse(dias[sabado.isoformat()]['sin_planificar'])  # tiene alternancia publicada
        self.assertTrue(dias[date(2026, 3, 8).isoformat()]['sin_planificar'])  # domingo sin sembrar
        self.assertNotIn(date(2026, 3, 4).isoformat(), dias)  # miércoles normal: no se pinta

    def test_endpoint_mes_responde_403_a_explorador_raso(self):
        self.client.force_login(self.a.user)
        resp = self.client.get('/turnos/api/reporte-mes/dias/', {'anio': 2026, 'mes': 3})
        self.assertEqual(resp.status_code, 403)

    def test_endpoint_mes_rechaza_parametros_invalidos(self):
        self.client.force_login(self.sup.user)
        for params in ({'anio': 'x', 'mes': 3}, {'anio': 2026, 'mes': 13}, {}):
            with self.subTest(params=params):
                resp = self.client.get('/turnos/api/reporte-mes/dias/', params)
                self.assertEqual(resp.status_code, 400)

    # ── 7c. El «con quién» del acuerdo ───────────────────────────────────────
    def _doblada_aprobada(self):
        """A cede su jornada de hoy a B, que la dobla. B trabaja, A descansa.

        Es el caso mínimo con las DOS caras del mismo trato: sirve para comprobar que las
        dos hojas de personas nombran al compañero y que la hoja de movimientos lo cuenta
        UNA sola vez.
        """
        from django.utils import timezone as _tz

        from solicitudes.models import DobladaDetalle
        from turnos.models import Turno

        sol = SolicitudCambio.objects.create(
            explorador_solicitante=self.a, explorador_receptor=self.b,
            tipo_cambio=self.tipo_dob, estado='aprobada', fecha_cambio_turno=FECHA,
            fecha_resolucion=_tz.now() - timedelta(days=3),
        )
        det = DobladaDetalle.objects.create(
            solicitud=sol, fecha_pago=FECHA + timedelta(days=7),
            tipo_cesion='cesion_completa', empleado_receptor=self.b,
        )
        # B queda doblando hoy (AM + PM); A sin turnos, que es como se aplica de verdad.
        for jor in (self.am, self.pm):
            Turno.objects.create(explorador=self.b, fecha=FECHA, jornada=jor,
                                 sala=self.sala, tipo_cambio='DOBLADA')
        det.snapshot_turnos_resultantes = {
            f'{self.b.id}:{FECHA.isoformat()}': [
                {'jornada_nombre': 'AM', 'sala_id': self.sala.id, 'tipo_cambio': 'DOBLADA'},
                {'jornada_nombre': 'PM', 'sala_id': self.sala.id, 'tipo_cambio': 'DOBLADA'},
            ],
            f'{self.a.id}:{FECHA.isoformat()}': [],   # el cedente queda libre
        }
        det.save(update_fields=['snapshot_turnos_resultantes'])
        return sol, det

    def test_las_dos_caras_del_acuerdo_nombran_al_companero(self):
        sol, det = self._doblada_aprobada()
        data = ReporteDiaService.reporte(FECHA)

        que_trabaja = _buscar(data, self.b.id)['acuerdo']
        que_descansa = _buscar(data, self.a.id)['acuerdo']
        self.assertEqual(que_trabaja['companero_nombre'], f'{self.a.nombre} {self.a.apellido}')
        self.assertEqual(que_trabaja['rol'], 'receptor')
        self.assertEqual(que_descansa['companero_nombre'], f'{self.b.nombre} {self.b.apellido}')
        self.assertEqual(que_descansa['rol'], 'solicitante')
        # Las dos caras hablan del MISMO trato y apuntan al MISMO día de devolución.
        self.assertEqual(que_trabaja['solicitud_id'], sol.id)
        self.assertEqual(que_descansa['solicitud_id'], sol.id)
        esperado = det.fecha_pago.strftime('%d/%m/%Y')
        self.assertEqual(que_trabaja['fecha_relacionada'], esperado)
        self.assertEqual(que_descansa['fecha_relacionada'], esperado)

    def test_excel_pone_el_companero_en_columna_propia_y_un_solo_movimiento(self):
        import openpyxl

        from turnos.api.views.reportes import ReporteDiaExcelView

        sol, det = self._doblada_aprobada()
        req = RequestFactory().get('/x?fecha=' + FECHA.isoformat())
        req.user = self.sup.user
        wb = openpyxl.load_workbook(io.BytesIO(ReporteDiaExcelView().get(req).content))

        def _fila_de(hoja, nombre):
            ws = wb[hoja]
            cols = {ws.cell(row=4, column=c).value: c for c in range(1, ws.max_column + 1)}
            for r in range(5, ws.max_row + 1):
                if ws.cell(row=r, column=cols['Nombre']).value == nombre:
                    return {k: ws.cell(row=r, column=c).value for k, c in cols.items()}
            return None

        # B dobla hoy: sale en las dos hojas de trabajo, nombrando a A.
        fila_b = _fila_de('Trabajan AM', self.b.nombre)
        self.assertEqual(fila_b['Con quién'], f'{self.a.nombre} {self.a.apellido}')
        self.assertEqual(fila_b['Tipo de cambio'], 'DOBLADA')
        self.assertEqual(fila_b['Su papel'], 'Recibió el cambio')
        self.assertEqual(fila_b['Fecha relacionada'], det.fecha_pago.strftime('%d/%m/%Y'))

        fila_a = _fila_de('Descansan', self.a.nombre)
        self.assertEqual(fila_a['Con quién'], f'{self.b.nombre} {self.b.apellido}')
        self.assertEqual(fila_a['Su papel'], 'Pidió el cambio')

        # Un acuerdo = UN movimiento, aunque aparezca en las dos personas.
        ws = wb['Cambios y permisos']
        cols = {ws.cell(row=4, column=c).value: c for c in range(1, ws.max_column + 1)}
        nums = [ws.cell(row=r, column=cols['N° solicitud']).value
                for r in range(5, ws.max_row + 1)]
        self.assertEqual([n for n in nums if n == sol.id], [sol.id])
        fila = next(r for r in range(5, ws.max_row + 1)
                    if ws.cell(row=r, column=cols['N° solicitud']).value == sol.id)
        self.assertEqual(ws.cell(row=fila, column=cols['Quién pidió']).value,
                         f'{self.a.nombre} {self.a.apellido}')
        self.assertEqual(ws.cell(row=fila, column=cols['Con quién']).value,
                         f'{self.b.nombre} {self.b.apellido}')
        # La frase junta los dos lados del trato.
        que_pasa = ws.cell(row=fila, column=cols['Qué pasa hoy']).value
        self.assertIn(f'{self.b.nombre} {self.b.apellido} trabaja', que_pasa)
        self.assertIn(f'{self.a.nombre} {self.a.apellido} descansa', que_pasa)

    def test_la_fecha_relacionada_nunca_es_el_dia_del_reporte(self):
        """
        En un CAMBIO DESCANSO los DOS lados figuran como 'cedio' (los dos ceden su día), así
        que decidir la contraparte por ese campo hacía que quien está en la fecha de pago se
        viera a sí mismo: "fecha relacionada = hoy", que no informa de nada. La otra fecha es
        la del par que NO es hoy.
        """
        from django.utils import timezone as _tz

        from solicitudes.models import DobladaDetalle

        tipo_cd = TipoSolicitudCambio.objects.create(
            nombre='CAMBIO DESCANSO', codigo_estrategia='CAMBIO DESCANSO', activo=True)
        cesion = FECHA - timedelta(days=3)
        sol = SolicitudCambio.objects.create(
            explorador_solicitante=self.a, explorador_receptor=self.b,
            tipo_cambio=tipo_cd, estado='aprobada', fecha_cambio_turno=cesion,
            fecha_resolucion=_tz.now() - timedelta(days=10),
        )
        DobladaDetalle.objects.create(solicitud=sol, fecha_pago=FECHA,
                                      tipo_cesion='cesion_completa', empleado_receptor=self.b)

        acuerdo = _buscar(ReporteDiaService.reporte(FECHA), self.a.id)['acuerdo']
        self.assertIsNotNone(acuerdo)
        self.assertEqual(acuerdo['fecha_relacionada'], cesion.strftime('%d/%m/%Y'))

    def test_el_reporte_no_gasta_mas_consultas_al_crecer_la_plantilla(self):
        """
        El «con quién» se resuelve en LOTE. Con ~400 exploradores, preguntarlo persona a
        persona serían ~400 tandas de consultas y el reporte se caería por tiempo: es toda
        la razón de ser de `AcuerdoPorDiaService.en_rango_multiple`.
        """
        from django.db import connection
        from django.test.utils import CaptureQueriesContext

        def _consultas():
            with CaptureQueriesContext(connection) as ctx:
                ReporteDiaService.reporte(FECHA)
            return len(ctx)

        base = _consultas()
        for n in range(20):
            u = User.objects.create_user(username=f'carga_{n}', password='x')
            e = Empleado.objects.create(user=u, nombre=f'Carga{n}', apellido='X',
                                        cedula=f'85{n:03d}', activo=True)
            AsignarJornadaExplorador.objects.create(
                explorador=e, jornada=self.am, fecha_inicio=date(2025, 1, 1))

        self.assertEqual(_consultas(), base,
                         'el número de consultas debe ser constante, no crecer por empleado')

    def test_excel_avisa_cuando_el_dia_no_tiene_movimientos(self):
        import openpyxl

        from turnos.api.views.reportes import ReporteDiaExcelView

        req = RequestFactory().get('/x?fecha=' + FECHA.isoformat())
        req.user = self.sup.user
        wb = openpyxl.load_workbook(io.BytesIO(ReporteDiaExcelView().get(req).content))
        self.assertIn('No hay cambios ni permisos', wb['Cambios y permisos']['A5'].value)

    def test_un_dia_sin_acuerdo_no_inventa_companero(self):
        """Un turno normal no tiene «con quién»: la columna queda en '—', no vacía ni falsa."""
        import openpyxl

        from turnos.api.views.reportes import ReporteDiaExcelView

        req = RequestFactory().get('/x?fecha=' + FECHA.isoformat())
        req.user = self.sup.user
        wb = openpyxl.load_workbook(io.BytesIO(ReporteDiaExcelView().get(req).content))
        ws = wb['Trabajan AM']
        cols = {ws.cell(row=4, column=c).value: c for c in range(1, ws.max_column + 1)}
        self.assertEqual(ws.cell(row=5, column=cols['Con quién']).value, '—')

    # ── 8. Permisos de la API ────────────────────────────────────────────────
    def test_api_sin_rol_supervisor_devuelve_403(self):
        self.client.force_login(self.a.user)  # explorador raso
        for url in ('/turnos/api/reporte-dia/', '/turnos/api/reporte-dia/excel/'):
            with self.subTest(url=url):
                resp = self.client.get(url, {'fecha': FECHA.isoformat()})
                self.assertEqual(resp.status_code, 403)
                self.assertEqual(resp.json()['error'], 'Sin permisos')

    def test_api_supervisor_fecha_invalida_o_ausente(self):
        self.client.force_login(self.sup.user)
        self.assertEqual(self.client.get('/turnos/api/reporte-dia/').status_code, 400)
        self.assertEqual(
            self.client.get('/turnos/api/reporte-dia/', {'fecha': 'ayer'}).status_code, 400)
        self.assertEqual(
            self.client.get('/turnos/api/reporte-dia/',
                            {'fecha': FECHA.isoformat()}).status_code, 200)

    def test_api_error_interno_no_filtra_la_excepcion(self):
        """El texto de la excepción puede llevar rutas o SQL: no debe salir al navegador."""
        from unittest.mock import patch
        self.client.force_login(self.sup.user)
        with patch('turnos.services.reporte_dia_service.ReporteDiaService.reporte',
                   side_effect=Exception('SELECT * FROM empleados_empleado -- /srv/app/secreto')):
            resp = self.client.get('/turnos/api/reporte-dia/', {'fecha': FECHA.isoformat()})
        self.assertEqual(resp.status_code, 500)
        self.assertNotIn('SELECT', resp.json()['error'])


class ReporteDiaOverrideFindeTest(TestCase):
    """
    El reporte del día debe leer la MISMA alternancia publicada que `estado_dia`.

    El servicio reimplementa las capas de `estado_dia` "en batch" y en esa copia miraba la
    alternancia cruda, así que un finde con asignación manual salía al revés en TODO el reporte:
    el grupo de la rotación aparecía trabajando y el que de verdad iba, descansando.
    """

    def setUp(self):
        from django.core.cache import cache
        cache.clear()
        self.am = Jornada.objects.create(nombre='AM', hora_inicio='06:00:00', hora_fin='14:00:00')
        self.pm = Jornada.objects.create(nombre='PM', hora_inicio='14:00:00', hora_fin='22:00:00')

        def _emp(username, ced, jor):
            u = User.objects.create_user(username=username, password='x')
            e = Empleado.objects.create(user=u, nombre=username, apellido='X', cedula=ced, activo=True)
            AsignarJornadaExplorador.objects.create(explorador=e, jornada=jor, fecha_inicio=date(2025, 1, 1))
            return e

        self.emp_am = _emp('ov_am', '9101', self.am)
        self.emp_pm = _emp('ov_pm', '9102', self.pm)

        # Un sábado futuro cualquiera.
        d = date(2026, 3, 7)
        while d.weekday() != 5:
            d += timedelta(days=1)
        self.sabado = d

    def _grupos(self):
        data = ReporteDiaService.reporte(self.sabado)
        trabajando = {e['id'] for e in data['trabajando']}
        return trabajando

    def _publicar(self, grupo):
        from django.core.cache import cache

        from turnos.models import AsignacionEspecialManual
        AsignacionEspecialManual.objects.update_or_create(
            fecha=self.sabado,
            defaults={'jornada_trabaja': self.am if grupo == 'AM' else self.pm,
                      'tipo': 'finde', 'activo': True})
        cache.clear()

    def test_el_reporte_sigue_la_alternancia_publicada(self):
        for grupo in ('AM', 'PM'):
            with self.subTest(grupo=grupo):
                self._publicar(grupo)
                esperado = self.emp_am if grupo == 'AM' else self.emp_pm
                el_otro = self.emp_pm if grupo == 'AM' else self.emp_am
                trabajando = self._grupos()
                self.assertIn(esperado.id, trabajando,
                              'debe trabajar el grupo que publicó el supervisor')
                self.assertNotIn(el_otro.id, trabajando,
                                 'el grupo contrario descansa ese día')

    def test_sin_alternancia_publicada_no_trabaja_nadie(self):
        """
        Sin fila no se inventa un grupo: nadie sale trabajando. Antes se caía a la fórmula
        y el reporte mostraba turnos de un año que el supervisor nunca publicó.
        """
        from turnos.services.asignacion_especial_service import AsignacionEspecialService
        self.assertIsNone(AsignacionEspecialService.grupo_trabaja(self.sabado))
        trabajando = self._grupos()
        self.assertNotIn(self.emp_am.id, trabajando)
        self.assertNotIn(self.emp_pm.id, trabajando)

    def test_sin_alternancia_se_distingue_de_un_descanso(self):
        """
        Que nadie trabaje NO puede verse igual que un día planificado en el que toca
        descansar: el reporte tiene que decir que el día está sin planificar.
        """
        data = ReporteDiaService.reporte(self.sabado)
        self.assertTrue(data['dia_info']['sin_planificar'])
        motivos = {e['motivo'] for e in data['descansando']}
        self.assertEqual(motivos, {'sin alternancia publicada'})

    def test_finde_publicado_no_queda_marcado_sin_planificar(self):
        self._publicar('AM')
        data = ReporteDiaService.reporte(self.sabado)
        self.assertFalse(data['dia_info']['sin_planificar'])
        self.assertEqual(
            [e['motivo'] for e in data['descansando']],
            ['descanso de fin de semana'],
        )
